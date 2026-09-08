#!/usr/bin/env python3
"""
Appian Enterprise Document Generator

Questo script genera un documento Excel (Enterprise Document) contenente
tutti gli oggetti di un'applicazione Appian esportata, con:
- Nome
- Descrizione
- Tipo di oggetto

Uso: python appian_enterprise_document.py <cartella_export_appian> [output.xlsx]
"""

import os
import sys
import xml.etree.ElementTree as ET
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def get_text(element, path, namespaces=None):
    """Estrae il testo da un elemento XML, gestendo vari formati."""
    if element is None:
        return ""
    
    # Prova con namespace
    if namespaces:
        for prefix, uri in namespaces.items():
            el = element.find(f".//{{{uri}}}{path}")
            if el is not None and el.text:
                return el.text.strip()
    
    # Prova senza namespace
    el = element.find(f".//{path}")
    if el is not None and el.text:
        return el.text.strip()
    
    # Prova con prefisso a:
    el = element.find(f".//a:{path}", {"a": "http://www.appian.com/ae/types/2009"})
    if el is not None and el.text:
        return el.text.strip()
    
    return ""


def get_name_from_stringmap(element):
    """Estrae il nome da una struttura string-map (usata nei process model)."""
    if element is None:
        return ""
    
    # Cerca value in tutti i figli (con o senza namespace)
    for elem in element.iter():
        if elem.tag.endswith('value') and elem.text:
            return elem.text.strip()
    
    # Fallback: cerca in string-map/pair/value senza namespace
    for pair in element.findall(".//pair"):
        value = pair.find("value")
        if value is not None and value.text:
            return value.text.strip()
    
    return ""


# Tag XML dell'oggetto -> etichetta leggibile.
# I nomi variano tra le versioni di Appian: teniamo entrambe le varianti note
# (es. decision/decisionTable, integration/outboundIntegration).
CONTENT_TYPES = {
    "rule": "Expression Rule",
    "interface": "Interface",
    "constant": "Constant",
    "rulesFolder": "Rules Folder",
    "document": "Document",
    "folder": "Folder",
    "documentFolder": "Document Folder",
    "integration": "Integration",
    "outboundIntegration": "Integration",
    "decision": "Decision",
    "decisionTable": "Decision",
    "report": "Report",
    "communityKnowledgeCenter": "Knowledge Center",
    "portal": "Portal",
    "processModelFolder": "Process Model Folder",
}

# Figli di contentHaul che sono metadati dell'export, non l'oggetto vero e proprio.
WRAPPER_CHILDREN = {"history", "roleMap", "versionUuid", "typedValue", "file"}


def local_tag(element):
    """Nome del tag senza namespace."""
    return element.tag.split("}")[-1] if "}" in element.tag else element.tag


def prettify_tag(tag):
    """camelCase -> 'Camel Case', per i tipi non ancora mappati."""
    out = tag[:1].upper()
    for ch in tag[1:]:
        out += f" {ch}" if ch.isupper() else ch
    return out


def parse_content_file(file_path):
    """
    Parsa un file XML dalla cartella content (o da qualsiasi cartella *Haul).

    L'oggetto è il figlio diretto della radice che non è un metadato di export.
    I tipi ignoti non vengono scartati: il tag stesso diventa l'etichetta, così
    l'inventario non perde oggetti quando Appian ne introduce di nuovi.
    """
    try:
        tree = ET.parse(file_path)
        root = tree.getroot()

        # Caso normale: <contentHaul><constant>...</constant><history/>...</contentHaul>
        for child in root:
            tag = local_tag(child)
            if tag in WRAPPER_CHILDREN:
                continue
            obj_type = CONTENT_TYPES.get(tag, prettify_tag(tag))
            name = get_text(child, "name")
            description = get_text(child, "description")
            if name:
                return {"name": name, "description": description, "type": obj_type}

        # Fallback per layout diversi: cerca il tag noto ovunque nell'albero
        for tag, obj_type in CONTENT_TYPES.items():
            element = root.find(f".//{tag}")
            if element is not None:
                name = get_text(element, "name")
                description = get_text(element, "description")
                if name:
                    return {"name": name, "description": description, "type": obj_type}

        return None
    except Exception as e:
        print(f"  Errore parsing {file_path}: {e}")
        return None


def parse_process_model_file(file_path):
    """Parsa un file XML di Process Model."""
    try:
        tree = ET.parse(file_path)
        root = tree.getroot()
        
        # Il nome è in pm/meta/name/string-map/pair/value (con namespace)
        # Prova con namespace
        meta = root.find(".//{http://www.appian.com/ae/types/2009}meta")
        if meta is None:
            meta = root.find(".//meta")
        
        if meta is not None:
            # Prova con namespace
            name_el = meta.find("{http://www.appian.com/ae/types/2009}name")
            if name_el is None:
                name_el = meta.find("name")
            name = get_name_from_stringmap(name_el)
            
            desc_el = meta.find("{http://www.appian.com/ae/types/2009}desc")
            if desc_el is None:
                desc_el = meta.find("desc")
            description = get_name_from_stringmap(desc_el)
            
            if name:
                return {"name": name, "description": description, "type": "Process Model"}
        
        return None
    except Exception as e:
        print(f"  Errore parsing {file_path}: {e}")
        return None


def parse_record_type_file(file_path):
    """Parsa un file XML di Record Type."""
    try:
        tree = ET.parse(file_path)
        root = tree.getroot()
        
        # Il nome è nell'attributo name del tag recordType
        record_type = root.find(".//{http://www.appian.com/ae/types/2009}recordType")
        if record_type is None:
            record_type = root.find(".//recordType")
        
        if record_type is not None:
            name = record_type.get("name", "") or record_type.get("{http://www.appian.com/ae/types/2009}name", "")
            description = get_text(record_type, "description")
            return {"name": name, "description": description, "type": "Record Type"}
        
        return None
    except Exception as e:
        print(f"  Errore parsing {file_path}: {e}")
        return None


def parse_datatype_file(file_path):
    """Parsa un file XSD di Data Type (CDT)."""
    try:
        tree = ET.parse(file_path)
        root = tree.getroot()
        
        # Il nome è nell'attributo name del complexType
        ns = {"xsd": "http://www.w3.org/2001/XMLSchema"}
        complex_type = root.find(".//xsd:complexType", ns)
        
        if complex_type is not None:
            name = complex_type.get("name", "")
            
            # La descrizione è in xsd:annotation/xsd:documentation
            doc = root.find(".//xsd:documentation", ns)
            description = doc.text.strip() if doc is not None and doc.text else ""
            
            return {"name": name, "description": description, "type": "Data Type (CDT)"}
        
        return None
    except Exception as e:
        print(f"  Errore parsing {file_path}: {e}")
        return None


def parse_group_file(file_path):
    """Parsa un file XML di Group."""
    try:
        tree = ET.parse(file_path)
        root = tree.getroot()
        
        group = root.find(".//group")
        if group is not None:
            name = get_text(group, "name")
            description = get_text(group, "description")
            return {"name": name, "description": description, "type": "Group"}
        
        return None
    except Exception as e:
        print(f"  Errore parsing {file_path}: {e}")
        return None


def parse_site_file(file_path):
    """Parsa un file XML di Site."""
    try:
        tree = ET.parse(file_path)
        root = tree.getroot()
        
        site = root.find(".//{http://www.appian.com/ae/types/2009}site")
        if site is None:
            site = root.find(".//site")
        
        if site is not None:
            name = site.get("name", "") or site.get("{http://www.appian.com/ae/types/2009}name", "")
            description = get_text(site, "description")
            return {"name": name, "description": description, "type": "Site"}
        
        return None
    except Exception as e:
        print(f"  Errore parsing {file_path}: {e}")
        return None


def parse_web_api_file(file_path):
    """Parsa un file XML di Web API."""
    try:
        tree = ET.parse(file_path)
        root = tree.getroot()
        
        web_api = root.find(".//{http://www.appian.com/ae/types/2009}webApi")
        if web_api is None:
            web_api = root.find(".//webApi")
        
        if web_api is not None:
            name = web_api.get("name", "") or web_api.get("{http://www.appian.com/ae/types/2009}name", "")
            description = get_text(web_api, "description")
            return {"name": name, "description": description, "type": "Web API"}
        
        return None
    except Exception as e:
        print(f"  Errore parsing {file_path}: {e}")
        return None


def parse_data_store_file(file_path):
    """Parsa un file XML di Data Store."""
    try:
        tree = ET.parse(file_path)
        root = tree.getroot()
        
        data_store = root.find(".//dataStore")
        if data_store is not None:
            name = get_text(data_store, "name")
            description = get_text(data_store, "description")
            return {"name": name, "description": description, "type": "Data Store"}
        
        return None
    except Exception as e:
        print(f"  Errore parsing {file_path}: {e}")
        return None


def parse_connected_system_file(file_path):
    """Parsa un file XML di Connected System."""
    try:
        tree = ET.parse(file_path)
        root = tree.getroot()
        
        cs = root.find(".//connectedSystem")
        if cs is not None:
            name = get_text(cs, "name")
            description = get_text(cs, "description")
            return {"name": name, "description": description, "type": "Connected System"}
        
        return None
    except Exception as e:
        print(f"  Errore parsing {file_path}: {e}")
        return None


def parse_application_file(file_path):
    """Parsa un file XML di Application."""
    try:
        tree = ET.parse(file_path)
        root = tree.getroot()
        
        app = root.find(".//application")
        if app is not None:
            name = get_text(app, "name")
            description = get_text(app, "description")
            return {"name": name, "description": description, "type": "Application"}
        
        return None
    except Exception as e:
        print(f"  Errore parsing {file_path}: {e}")
        return None


# Cartelle note dell'export -> (parser, etichetta). Le altre vengono comunque
# attraversate dallo scanner con il parser generico.
FOLDER_PARSERS = [
    ("application", parse_application_file, "Application"),
    ("connectedSystem", parse_connected_system_file, "Connected System"),
    ("content", parse_content_file, "Content"),
    ("dataStore", parse_data_store_file, "Data Store"),
    ("datatype", parse_datatype_file, "Data Type"),
    ("group", parse_group_file, "Group"),
    ("processModel", parse_process_model_file, "Process Model"),
    ("recordType", parse_record_type_file, "Record Type"),
    ("site", parse_site_file, "Site"),
    ("webApi", parse_web_api_file, "Web API"),
]

# Cartelle di servizio dell'archivio, non contengono oggetti dell'applicazione.
SKIP_FOLDERS = {"META-INF", "__MACOSX"}


def scan_appian_export_with_stats(export_path, verbose=True):
    """
    Scansiona l'export e restituisce (oggetti, statistiche).

    Attraversa *tutte* le cartelle dell'export, non solo quelle note: le cartelle
    non mappate vengono lette con il parser generico. Le statistiche riportano
    quanti file XML sono stati esaminati e quali non hanno prodotto un oggetto,
    così un tipo non gestito si vede subito invece di sparire in silenzio.
    """
    export_path = Path(export_path)
    all_objects = []
    scanned = 0
    unparsed = []

    if verbose:
        print("\nScansione export Appian in corso...\n")

    parsers = {name: func for name, func, _ in FOLDER_PARSERS}
    labels = {name: label for name, _, label in FOLDER_PARSERS}

    # Cartelle note prima (ordine stabile), poi tutte le altre presenti
    known = [name for name, _, _ in FOLDER_PARSERS]
    extra = sorted(
        d.name for d in export_path.iterdir()
        if d.is_dir() and d.name not in known and d.name not in SKIP_FOLDERS
    )

    for folder_name in known + extra:
        folder = export_path / folder_name
        label = labels.get(folder_name, folder_name)

        if not folder.exists():
            if verbose:
                print(f"  [--] {label}: cartella non presente")
            continue

        parser_func = parsers.get(folder_name, parse_content_file)
        count = 0
        missed = 0

        for file_path in sorted(folder.iterdir()):
            if file_path.name.startswith("._"):
                continue
            if file_path.suffix.lower() not in (".xml", ".xsd"):
                continue

            scanned += 1
            result = parser_func(file_path)
            if result and result.get("name"):
                all_objects.append(result)
                count += 1
            else:
                missed += 1
                unparsed.append({
                    "folder": folder_name,
                    "file": file_path.name,
                    "tags": describe_file(file_path),
                })

        if verbose:
            suffix = f"  ({missed} non riconosciuti)" if missed else ""
            note = "" if folder_name in labels else "  [cartella non mappata]"
            print(f"  [OK] {label}: {count} oggetti trovati{suffix}{note}")

    stats = {
        "files_scanned": scanned,
        "objects_found": len(all_objects),
        "unparsed": unparsed,
    }
    return all_objects, stats


def describe_file(file_path):
    """Tag dei figli non-metadato di un file, per capire cosa non è stato riconosciuto."""
    try:
        root = ET.parse(file_path).getroot()
    except Exception as e:
        return f"<illeggibile: {type(e).__name__}>"
    tags = [local_tag(c) for c in root if local_tag(c) not in WRAPPER_CHILDREN]
    return ", ".join(tags) if tags else f"<{local_tag(root)} senza oggetto>"


def scan_appian_export(export_path):
    """Scansiona l'export e restituisce la lista degli oggetti."""
    objects, _ = scan_appian_export_with_stats(export_path)
    return objects


def format_coverage(stats):
    """Riepilogo testuale del controllo di copertura."""
    lines = []
    scanned = stats["files_scanned"]
    found = stats["objects_found"]
    unparsed = stats["unparsed"]

    lines.append(f"File XML esaminati:  {scanned}")
    lines.append(f"Oggetti riconosciuti: {found}")

    if not unparsed:
        lines.append("Copertura: 100% - nessun file scartato.")
        return "\n".join(lines)

    pct = 100.0 * found / scanned if scanned else 0.0
    lines.append(f"Copertura: {pct:.1f}% - {len(unparsed)} file senza oggetto:")

    by_tag = {}
    for item in unparsed:
        by_tag.setdefault((item["folder"], item["tags"]), []).append(item["file"])
    for (folder, tags), files in sorted(by_tag.items()):
        lines.append(f"  {len(files):4}  {folder}/  ->  {tags}")
        for f in files[:3]:
            lines.append(f"          {f}")
        if len(files) > 3:
            lines.append(f"          ... e altri {len(files) - 3}")
    return "\n".join(lines)


def create_excel(objects, output_path, app_name="Appian Application"):
    """Crea il file Excel con l'Enterprise Document."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Enterprise Document"
    
    # Stili
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=12)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Header
    headers = ["Nome", "Descrizione", "Tipo Oggetto"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    
    # Ordina gli oggetti per tipo, poi per nome
    objects_sorted = sorted(objects, key=lambda x: (x["type"], x["name"]))
    
    # Dati
    for row, obj in enumerate(objects_sorted, 2):
        ws.cell(row=row, column=1, value=obj["name"]).border = border
        ws.cell(row=row, column=2, value=obj["description"]).border = border
        ws.cell(row=row, column=3, value=obj["type"]).border = border
    
    # Larghezza colonne
    ws.column_dimensions['A'].width = 50
    ws.column_dimensions['B'].width = 60
    ws.column_dimensions['C'].width = 25
    
    # Freeze prima riga
    ws.freeze_panes = 'A2'
    
    # Salva
    wb.save(output_path)
    print(f"\nFile Excel salvato: {output_path}")


def main():
    if len(sys.argv) < 2:
        print("Uso: python appian_enterprise_document.py <cartella_export_appian> [output.xlsx]")
        print("\nEsempio:")
        print("  python appian_enterprise_document.py './export'")
        print("  python appian_enterprise_document.py './export' 'enterprise_doc.xlsx'")
        sys.exit(1)
    
    export_path = sys.argv[1]
    
    if not os.path.isdir(export_path):
        print(f"Errore: '{export_path}' non è una cartella valida")
        sys.exit(1)
    
    # Nome output
    if len(sys.argv) >= 3:
        output_file = sys.argv[2]
    else:
        folder_name = os.path.basename(os.path.normpath(export_path))
        output_file = f"{folder_name}_Enterprise_Document.xlsx"
    
    # Scansiona e crea Excel
    objects, stats = scan_appian_export_with_stats(export_path)
    
    if objects:
        print(f"\nTotale oggetti trovati: {len(objects)}")
        create_excel(objects, output_file, os.path.basename(export_path))
        
        # Riepilogo per tipo
        print("\nRiepilogo per tipo:")
        type_counts = {}
        for obj in objects:
            t = obj["type"]
            type_counts[t] = type_counts.get(t, 0) + 1
        
        for t, count in sorted(type_counts.items()):
            print(f"   {t}: {count}")
        
        # Controllo di copertura
        print("\nControllo copertura:")
        print(format_coverage(stats))
    else:
        print("\nNessun oggetto trovato nell'export")
        print(format_coverage(stats))


if __name__ == "__main__":
    main()
